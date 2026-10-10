# app/api/reports.py - 基于本体的报告/PPT 生成（R8：工具层扩展）
# POST /api/projects/{pid}/reports/generate
#   kind=report：采集本体结构统计 + 样本实体 → LLM 生成 markdown 报告 → 附 md/docx 下载
#   kind=ppt：LLM 生成 slide 大纲 JSON → python-pptx 构建简单演示文稿 → 附 pptx 下载
# 数据全部来自本项目 Entity/Relation/UploadedDocument 实时统计，无硬编码模板词。
from __future__ import annotations

import io
import json
import logging
import re
import time
import uuid
from collections import Counter
from datetime import datetime

from fastapi import APIRouter, Depends, File, UploadFile
from lxml import etree
from pptx.oxml.ns import qn
from pydantic import BaseModel, Field
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.api.qa import _ensure_qa_access
from app.core.deps import ensure_module, get_current_user, get_db, require_any_module
from app.core.exceptions import APIError
from app.infrastructure.database import Entity, Project, ReasoningResult, Relation, UploadedDocument, User
from app.infrastructure.llm_client import LLMClient
from app.infrastructure.minio_client import get_minio_client
from app.core.config import settings
from app.adapters.provider import build_legacy_llm_config
from app.services.audit_service import log_action

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/projects", tags=["reports"])

MAX_SLIDES_LIMIT = 15


class ReportRequest(BaseModel):
    kind: str = Field(pattern="^(report|ppt)$")
    topic: str = Field(min_length=1, max_length=300)  # 分析主题（必填，面向业务）
    max_slides: int = Field(default=8, ge=3, le=MAX_SLIDES_LIMIT)
    use_template: bool = False  # kind=ppt 时使用已上传的自定义模板
    include_inferred: bool = False  # 附加「推理衍生事实」附录（语义推理产物，不进正文）


def _template_key(project_id: int) -> str:
    return f"ppt-templates/{project_id}/template.pptx"


def _template_meta_key(project_id: int) -> str:
    return f"ppt-templates/{project_id}/meta.json"


@router.post("/{project_id}/reports/ppt-template", dependencies=[Depends(require_any_module("report", "ppt"))])
async def upload_ppt_template(project_id: int, file: UploadFile = File(...),
                              db: Session = Depends(get_db),
                              user: User = Depends(get_current_user)):
    """上传 PPT 模板（.pptx，占位符协议：封面 {{title}}/{{subtitle}}/{{date}}/{{project}}，
    内容页 {{slide_title}}+{{bullets}}，结尾页 {{closing}}；覆盖式保存）。"""
    _ensure_qa_access(db, project_id, user)
    data = await file.read()
    if len(data) > 20 * 1024 * 1024:
        raise APIError("模板文件不能超过 20MB", code="PPT_TEMPLATE_TOO_LARGE", http_status=400)
    try:
        from pptx import Presentation
        Presentation(io.BytesIO(data))
    except Exception:
        raise APIError("不是有效的 .pptx 文件", code="PPT_TEMPLATE_INVALID", http_status=400)
    get_minio_client().put_bytes(
        settings.MINIO_BUCKET_EXPORTS, _template_key(project_id), data,
        content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation")
    get_minio_client().put_bytes(
        settings.MINIO_BUCKET_EXPORTS, _template_meta_key(project_id),
        json.dumps({"filename": file.filename}, ensure_ascii=False).encode(),
        content_type="application/json")
    log_action(db, user.id, "report.ppt_template_upload", resource_type="project",
               resource_id=str(project_id), detail={"filename": file.filename})
    db.commit()
    return {"ok": True, "filename": file.filename, "size": len(data)}


@router.get("/{project_id}/reports/ppt-template", dependencies=[Depends(require_any_module("report", "ppt"))])
def get_ppt_template(project_id: int, db: Session = Depends(get_db),
                     user: User = Depends(get_current_user)):
    _ensure_qa_access(db, project_id, user)
    stat = get_minio_client().stat_object(settings.MINIO_BUCKET_EXPORTS, _template_key(project_id))
    if stat is None:
        return {"exists": False}
    filename = f"template_{project_id}.pptx"
    try:
        raw = get_minio_client().get_bytes(settings.MINIO_BUCKET_EXPORTS, _template_meta_key(project_id))
        filename = json.loads(raw).get("filename", filename)
    except Exception:
        pass
    return {"exists": True, "filename": filename, "size": stat.get("size")}


@router.get("/{project_id}/reports/ppt-template/download", dependencies=[Depends(require_any_module("report", "ppt"))])
def download_ppt_template(project_id: int, db: Session = Depends(get_db),
                          user: User = Depends(get_current_user)):
    """下载当前已上传的模板（预签名 URL，1 小时有效）。"""
    _ensure_qa_access(db, project_id, user)
    mc = get_minio_client()
    if mc.stat_object(settings.MINIO_BUCKET_EXPORTS, _template_key(project_id)) is None:
        raise APIError("尚未上传模板", code="PPT_TEMPLATE_MISSING", http_status=400)
    url = mc.get_presigned_url(settings.MINIO_BUCKET_EXPORTS, _template_key(project_id))
    return {"url": url}


@router.get("/{project_id}/reports/ppt-template/sample", dependencies=[Depends(require_any_module("report", "ppt"))])
def download_sample_template(project_id: int, db: Session = Depends(get_db),
                             user: User = Depends(get_current_user)):
    """下载内置示例模板（内含占位符协议说明页，生成时说明页自动删除）。"""
    _ensure_qa_access(db, project_id, user)
    return {"url": _ensure_sample_template()}


@router.delete("/{project_id}/reports/ppt-template", dependencies=[Depends(require_any_module("report", "ppt"))])
def delete_ppt_template(project_id: int, db: Session = Depends(get_db),
                        user: User = Depends(get_current_user)):
    _ensure_qa_access(db, project_id, user)
    for key in (_template_key(project_id), _template_meta_key(project_id)):
        get_minio_client().remove_object(settings.MINIO_BUCKET_EXPORTS, key)
    return {"ok": True}


# ---------------------------------------------------------------- 上下文采集

def _collect_context(db: Session, project_id: int, topic: str | None = None) -> dict:
    """实时统计：类分布 / 关系谓词分布 / 每类样本实体 / 文档清单；
    给定 topic 时追加主题证据（直接匹配的实体及其关系），供主题分析报告引用。"""
    ents = db.query(Entity).filter(Entity.project_id == project_id).all()
    rels = (db.query(Relation)
            .filter(Relation.project_id == project_id, Relation.is_class_edge.is_(False)).all())

    class_counter = Counter(e.class_label for e in ents)
    pred_counter = Counter(r.predicate for r in rels)

    by_class: dict[str, list[Entity]] = {}
    for e in ents:
        by_class.setdefault(e.class_label, []).append(e)

    top_classes = class_counter.most_common(15)
    samples = []
    for cls, cnt in top_classes:
        sample_ents = by_class[cls][:6]
        samples.append({
            "class": cls,
            "count": cnt,
            "entities": [
                {
                    "label": e.label,
                    "props": {k: v for k, v in (e.props or {}).items()
                              if isinstance(v, (str, int, float, bool))} if e.props else {},
                }
                for e in sample_ents
            ],
        })

    docs = (db.query(UploadedDocument)
            .filter(UploadedDocument.project_id == project_id, UploadedDocument.deleted_at.is_(None))
            .limit(30).all())

    ctx = {
        "relation_total": len(rels),
        "entity_total": len(ents),
        "top_relations": [{"predicate": p, "count": c} for p, c in pred_counter.most_common(15)],
        "classes": samples,
        "documents": [d.filename for d in docs],
    }

    # ── 主题证据：label/别名 直接包含主题关键词的实体 + 其出入关系
    if topic:
        kw = topic.strip()
        hit = [e for e in ents
               if e.is_class_node is False and e.status != "rejected"
               and (kw in (e.label or "") or kw in (e.class_label or "")
                    or any(kw in str(a) for a in (e.aliases or [])))]
        # 命中不足时放宽为主题前两字片段（中文无分词器的兜底）
        if len(hit) < 5 and len(kw) >= 2:
            frag = kw[:2]
            hit += [e for e in ents
                    if e.is_class_node is False and e.status != "rejected"
                    and e not in hit and (frag in (e.label or "") or frag in (e.class_label or ""))]
        hit = hit[:30]
        hit_ids = {e.id for e in hit}
        topic_rels = [
            r for r in rels if r.subject_id in hit_ids or r.object_id in hit_ids
        ][:60]
        label_of = {e.id: e.label for e in ents}
        ctx["topic"] = kw
        ctx["topic_entities"] = [
            {
                "label": e.label, "class_label": e.class_label,
                "props": {k: v for k, v in (e.props or {}).items()
                          if isinstance(v, (str, int, float, bool))} if e.props else {},
            }
            for e in hit
        ]
        ctx["topic_relations"] = [
            {"subject": label_of.get(r.subject_id, str(r.subject_id)),
             "predicate": r.predicate,
             "object": label_of.get(r.object_id, str(r.object_id))}
            for r in topic_rels
        ]
    return ctx


def _context_text(ctx: dict) -> str:
    lines: list[str] = []
    # 主题证据优先（主题分析报告的核心材料）
    if ctx.get("topic"):
        lines.append(f"分析主题：{ctx['topic']}")
        if ctx.get("topic_entities"):
            lines.append("- 主题相关实体（知识库中与主题直接匹配）：")
            for e in ctx["topic_entities"][:30]:
                props = "；".join(f"{k}={v}" for k, v in list(e["props"].items())[:6])
                lines.append(f"  · {e['label']}（{e['class_label']}）" + (f"：{props}" if props else ""))
        else:
            lines.append("- 知识库中与主题直接匹配的实体较少（材料以全局核心实体为主，报告须说明知识覆盖边界）")
        if ctx.get("topic_relations"):
            lines.append("- 主题相关关系（证据链）：")
            for r in ctx["topic_relations"][:40]:
                lines.append(f"  · {r['subject']} —[{r['predicate']}]\u2192 {r['object']}")
    lines.append("- 实体总数：" + str(ctx["entity_total"]))
    lines.append("- 关系总数（实例级）：" + str(ctx["relation_total"]))
    lines.append("- 文档清单：" + (", ".join(ctx["documents"]) if ctx["documents"] else "无"))
    lines.append("- 类分布与样本实体（全局覆盖范围）：")
    for c in ctx["classes"]:
        names = "；".join(f"{e['label']}" + (f"（{json.dumps(e['props'], ensure_ascii=False)}）"
                                             if e["props"] else "") for e in c["entities"])
        lines.append(f"  · {c['class']}（{c['count']} 个实例）：{names}")
    lines.append("- 主要关系谓词：" + (", ".join(f"{r['predicate']}({r['count']})" for r in ctx["top_relations"]) or "无"))
    return "\n".join(lines)


# ---------------------------------------------------------------- 推理附录

MAX_APPENDIX_ROWS = 50  # 附录表格至多展示的推理衍生三元组条数


def _inference_rows(db: Session, project_id: int) -> list[ReasoningResult]:
    """最新推理批次的全部衍生三元组（latest-only 语义，一次 run 覆盖前次）。"""
    batch = (db.query(ReasoningResult.batch_id)
             .filter(ReasoningResult.project_id == project_id)
             .order_by(ReasoningResult.id.desc()).limit(1).first())
    if not batch:
        return []
    return (db.query(ReasoningResult)
            .filter(ReasoningResult.project_id == project_id,
                    ReasoningResult.batch_id == batch.batch_id)
            .order_by(ReasoningResult.id.asc()).all())


def _inference_origin(r: ReasoningResult) -> str:
    return f"规则「{r.rule_name}」" if r.rule_name else f"蕴含推理（{r.source}）"


def _inference_cell(v: str) -> str:
    # markdown 表格防串列：竖线替换为斜杠
    return (v or "").replace("|", "／")


def _inference_appendix_md(rows: list[ReasoningResult]) -> str:
    """附录 markdown（确定性拼接，不经 LLM）：推理产物不得冒充事实——
    正文生成完全不使用本节数据，仅在报告末尾以表格披露并注明非原始事实记载。"""
    lines = [
        "## 附录：推理衍生事实（语义推理产物，非原始事实记载）",
        "",
        "以下结论由系统基于本体公理与自定义规则自动推导，"
        "不是文档原文记载，仅供延伸参考；正文分析未使用本节内容：",
        "",
        "| # | 主体 | 关系 | 客体 | 推导方式 |",
        "| --- | --- | --- | --- | --- |",
    ]
    for i, r in enumerate(rows[:MAX_APPENDIX_ROWS], 1):
        lines.append(f"| {i} | {_inference_cell(r.subject_label)} | {_inference_cell(r.predicate_label)} "
                     f"| {_inference_cell(r.object_label)} | {_inference_origin(r)} |")
    if len(rows) > MAX_APPENDIX_ROWS:
        lines.append("")
        lines.append(f"> 共 {len(rows)} 条推导结果，本表至多展示 {MAX_APPENDIX_ROWS} 条。")
    return "\n".join(lines)


def _inference_appendix_slide(rows: list[ReasoningResult]) -> dict:
    """PPT 附录页大纲（确定性追加，置于结尾页之前）：标题即声明产物性质。"""
    bullets = ["以下为语义推理产物（非原始事实记载），仅供延伸参考"]
    for r in rows[:4]:
        bullets.append(f"{r.subject_label} —[{r.predicate_label}]→ {r.object_label}（{_inference_origin(r)}）")
    if len(rows) > 4:
        bullets.append(f"共 {len(rows)} 条推导结果，此处展示前 4 条")
    return {"title": "附录：推理衍生事实", "bullets": bullets}


# ---------------------------------------------------------------- LLM 生成

def _llm_call(llm: LLMClient, system: str, user: str) -> str:
    result = llm.call_llm_text(system, user, timeout=300)
    # call_llm_text 返回 dict（含 content/usage）或 str，做兼容展开
    if isinstance(result, dict):
        content = result.get("content") or result.get("text") or ""
    else:
        content = str(result)
    if not content.strip():
        raise APIError("模型未返回内容，请稍后重试", code="REPORT_EMPTY", http_status=502)
    return content


REPORT_SYSTEM = (
    "你是一名资深行业分析师。用户给你一个分析主题，以及一份来自企业知识图谱的证据材料"
    "（主题相关实体及其属性、实体间关系、知识库全局覆盖范围）。"
    "请撰写一份面向业务读者的中文主题分析报告（markdown）。要求："
    "1. 报告标题体现主题；正文完全围绕主题展开：开头给出执行摘要与核心结论，"
    "中间分 2~4 个分析维度逐层展开（每个论点都要有具体事实支撑），"
    "结尾给出业务建议与结论；"
    "2. 所有事实必须来自给定的证据材料，不得编造；"
    "3. 行文规范（重要）：像一份正式的行业研究报告那样直接陈述事实，"
    "例如『该产品为封闭式固定收益类产品』『管理人通过中国理财网披露净值』，"
    "严禁使用『材料/证据/知识库/图谱/记载/显示/提到/列示/披露于材料』等暴露来源的元话语，"
    "严禁出现『材料中』『材料记载』『根据材料』『材料说明』『上述材料』等表述；"
    "如需交代分析范围，全文最多在开头用一句『本报告基于该产品说明书披露信息撰写』；"
    "表格列头用业务词（如『风险类别/对应安排』），不得用『材料中的依据』这类说法；"
    "4. 语言面向业务读者（如管理层、产品与风控人员），"
    "禁止把『实体类/关系谓词/instance_of/图谱统计』等建模术语当作章节主体；"
    "5. 关键对比信息用 markdown 表格呈现；篇幅 1000~1800 字；"
    "6. 加粗保持克制：仅用于各章节的核心结论短语（每段至多 1 处）；"
    "编号/列表项如果要表达小节标题，整项加粗或整项不加粗，不得只在项内局部加粗几个字；"
    "7. 若证据不足以完全支撑主题，在合适位置客观说明分析边界"
    "（如『相关条款说明书未进一步展开』），同样不得用『材料』一词。"
)

PPT_SYSTEM = (
    "你是一名演示文稿策划专家。用户给你一个分析主题，以及一份来自企业知识图谱的证据材料"
    "（主题相关实体及其属性、实体间关系、知识库全局覆盖范围）。"
    "请输出一个面向业务读者的中文演示文稿大纲。只输出 JSON，不要 markdown 代码块。"
    'JSON 格式：{"title": "演示文稿标题（体现主题）", "subtitle": "副标题", "slides": '
    '[{"title": "页标题", "bullets": ["要点1", "要点2"]}]}。'
    "要求：5~10 页，每页 3~5 条要点（每条不超过 40 字）；"
    "结构围绕主题：背景与问题 → 核心发现（2~4 页，每条要点有具体事实支撑）→ "
    "影响与建议 → 结论；"
    "行文规范（重要）：像正式的行业研究演示那样直接陈述事实，"
    "严禁使用『材料/证据/知识库/图谱/记载/显示/提到/列示』等暴露来源的元话语，"
    "严禁出现『材料中』『材料记载』『根据材料』等表述；"
    "页标题与要点面向业务读者，不得使用『实体类/关系谓词/instance_of』等建模术语；"
    "不得编造给定事实之外的内容。"
)


def _parse_ppt_outline(raw: str) -> dict:
    text = raw.strip()
    m = re.search(r"```(?:json)?\s*(.*?)```", text, re.S)
    if m:
        text = m.group(1).strip()
    start, end = text.find("{"), text.rfind("}")
    if start >= 0 and end > start:
        text = text[start:end + 1]
    outline = json.loads(text)
    if not isinstance(outline.get("slides"), list) or not outline["slides"]:
        raise APIError("模型未生成有效的大纲", code="REPORT_BAD_OUTLINE", http_status=502)
    outline["slides"] = outline["slides"][:MAX_SLIDES_LIMIT]
    return outline


# ---------------------------------------------------------------- 文档导出

def _markdown_to_docx(md_text: str, *, title: str = "本体分析报告", project_name: str = "",
                      entity_total: int = 0, relation_total: int = 0) -> bytes:
    """md→docx（R8 排版升级）：封面页 + 真 Word 表格 + 中文字体/1.3 行距/首行缩进 + 页眉页脚页码。"""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    ACCENT = RGBColor(0x1F, 0x4E, 0x79)          # 标题深商务蓝
    HEAD_FILL = "DCE6F1"                          # 表头底纹
    GRAY = RGBColor(0x7F, 0x7F, 0x7F)
    EA_FONT = "微软雅黑"
    LS_FONT = "Calibri"

    doc = Document()

    # ── 全局样式：正文与三级标题（中西文字体、行距、间距、颜色）
    normal = doc.styles["Normal"]
    normal.font.name = LS_FONT
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), EA_FONT)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(6)
    for lvl, size in ((1, 18), (2, 14), (3, 12), (4, 11)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = LS_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = ACCENT
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), EA_FONT)
        st.paragraph_format.space_before = Pt(16 if lvl <= 2 else 10)
        st.paragraph_format.space_after = Pt(6)

    def _add_rich(p, text: str, *, size: float | None = None, color: RGBColor | None = None,
                  base_bold: bool = False, italic: bool = False):
        """内联 **bold** / `code` → 富文本 runs；bold 仅显式设置 True，False 交给样式继承
        （否则 run 级 bold=False 会覆盖 Heading 样式的加粗，导致标题粗细不一）。"""
        for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
            if not part:
                continue
            r_bold, r_mono = base_bold, False
            content = part
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                content, r_bold = part[2:-2], True
            elif part.startswith("`") and part.endswith("`") and len(part) > 2:
                content, r_mono = part[1:-1], True
            r = p.add_run(content)
            if r_bold:
                r.bold = True
            if italic:
                r.italic = True
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
            if r_mono:
                r.font.name = "Consolas"
                r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), EA_FONT)
        return p

    def _shade(cell, fill: str):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), fill)
        tcPr.append(shd)

    def _add_table(header: list[str], rows: list[list[str]]):
        cols = len(header)
        t = doc.add_table(rows=len(rows) + 1, cols=cols)
        try:
            t.style = doc.styles["Table Grid"]
        except KeyError:
            pass
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = True
        # 表头行：底纹 + 加粗 + 跨页重复（tableHeader）
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trPr.append(th)
        for j, text in enumerate(header):
            cell = t.rows[0].cells[j]
            _shade(cell, HEAD_FILL)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            _add_rich(p, text, size=10, color=ACCENT, base_bold=True)
        for i, row in enumerate(rows, start=1):
            for j in range(cols):
                p = t.rows[i].cells[j].paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                _add_rich(p, row[j] if j < len(row) else "", size=10)
        gap = doc.add_paragraph()  # 表后留隙（防相邻表格粘连）
        gap.paragraph_format.space_after = Pt(4)
        for r in gap.runs:
            r.font.size = Pt(2)

    def _field(run, instr: str):
        f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
        f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
        run._r.append(f1); run._r.append(it); run._r.append(f2)

    # ── 页眉（报告名）+ 页脚（第 X 页 / 共 Y 页）
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run(title)
    hr.font.size = Pt(9)
    hr.font.color.rgb = GRAY
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("第 ")
    _field(fp.add_run(), "PAGE")
    fp.add_run(" 页 / 共 ")
    _field(fp.add_run(), "NUMPAGES")
    fp.add_run(" 页")
    for r in fp.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = GRAY

    # ── 封面页：留白 + 居中主标题 + 项目/统计/时间元信息 → 分页
    for _ in range(7):
        doc.add_paragraph()
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_rich(tp, title, size=26, color=ACCENT, base_bold=True)
    now = datetime.now()
    for line in (f"项目名称：{project_name or '—'}",
                 f"实体 {entity_total} 个 · 关系 {relation_total} 条",
                 f"生成时间：{now.strftime('%Y 年 %m 月 %d 日')}"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        _add_rich(p, line, size=12, color=GRAY)
    doc.add_page_break()

    # ── 正文：md 逐行状态机（表格块收集为真表格）
    lines = md_text.splitlines()
    i = 0
    first_h1_skipped = False
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue
        # markdown 表格：当前行以 | 开头且下一行是 |---| 分隔行 → 收集整块
        if stripped.startswith("|") and i + 1 < len(lines) \
                and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1] or ""):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.match(r"^:?-+:?$", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            _add_table(header, rows)
            continue
        if stripped.startswith("#### "):
            doc.add_heading(stripped[5:].strip(), level=4)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("# "):
            if first_h1_skipped:  # 首个 H1 已进封面
                doc.add_heading(stripped[2:].strip(), level=1)
            first_h1_skipped = True
        elif re.match(r"^[-*]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            _add_rich(p, re.sub(r"^[-*]\s+", "", stripped))
        elif re.match(r"^\d+[.、)]\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            _add_rich(p, re.sub(r"^\d+[.、)]\s+", "", stripped))
        elif stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _add_rich(p, stripped.lstrip("> ").strip(), italic=True, color=GRAY)
        elif re.match(r"^(-{3,}|\*{3,})$", stripped):
            pass  # 分隔线跳过（章节间距已足够）
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Pt(22)  # 中文正文首行缩进两字符
            _add_rich(p, stripped)
        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_pptx(outline: dict, project_name: str = "") -> bytes:
    """主题演示文稿（R8 排版升级）：16:9 + 深蓝封面/装饰几何 + 内容页竖条标题/
    彩色项目符号/页脚页码 + 结尾页；全文微软雅黑。"""
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    DEEP = RGBColor(0x1B, 0x3C, 0x6E)     # 封面/结尾底色
    DEEP2 = RGBColor(0x25, 0x50, 0x8F)    # 封面装饰圆
    ACCENT = RGBColor(0x5B, 0x8D, 0xEF)   # 平台蓝：竖条/符号/短线
    HEAD = RGBColor(0x1F, 0x4E, 0x79)     # 内容页标题
    TEXT = RGBColor(0x33, 0x3A, 0x45)     # 正文
    GRAY = RGBColor(0x8A, 0x94, 0xA6)     # 页脚
    PALE = RGBColor(0xBF, 0xD4, 0xF5)     # 封面副文字
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    FONT = "微软雅黑"

    def _styled(run, size: float, color: RGBColor, bold: bool = False):
        f = run.font
        f.size = Pt(size)
        f.color.rgb = color
        f.bold = bold
        f.name = FONT
        _apply_ea_font(run._r.get_or_add_rPr(), FONT)  # 中文字形走 eastAsia

    def _rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
        sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sp.fill.solid()
        sp.fill.fore_color.rgb = color
        sp.line.fill.background()
        sp.shadow.inherit = False
        return sp

    def _text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        """lines: [(text, size, color, bold, space_after)]"""
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for i, (text, size, color, bold, space_after) in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = Pt(space_after)
            p.line_spacing = 1.2
            r = p.add_run()
            r.text = text
            _styled(r, size, color, bold)
        return tb

    def _clean(v) -> str:
        return str(v).replace("**", "").replace("`", "").strip()

    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    now = datetime.now().strftime("%Y-%m-%d")
    footer = f"OntologySystem · {project_name or '本体分析'} · {now}"

    # ── 封面：深蓝底 + 装饰圆 + 竖条 + 主/副标题 + 落款
    s = prs.slides.add_slide(blank)
    _rect(s, 0, 0, 13.333, 7.5, DEEP)
    _rect(s, 9.9, -2.3, 5.8, 5.8, DEEP2, MSO_SHAPE.OVAL)
    _rect(s, -1.5, 5.5, 3.6, 3.6, DEEP2, MSO_SHAPE.OVAL)
    _rect(s, 0.95, 2.62, 0.10, 1.62, ACCENT)
    _text(s, 1.35, 2.42, 10.7, 2.1,
          [(_clean(outline.get("title") or "主题分析报告"), 33, WHITE, True, 6)])
    if outline.get("subtitle"):
        _text(s, 1.35, 4.62, 10.6, 0.9,
              [(_clean(outline["subtitle"]), 16, PALE, False, 0)])
    _text(s, 1.35, 6.74, 10.5, 0.5, [(footer, 11, PALE, False, 0)])

    # ── 内容页：竖条标题 + 短线 + 彩色项目符号正文 + 页脚/页码
    for idx, item in enumerate(outline["slides"], start=1):
        s = prs.slides.add_slide(blank)
        _rect(s, 0.55, 0.50, 0.10, 0.48, ACCENT)
        _text(s, 0.82, 0.38, 11.8, 0.75,
              [(_clean(item.get("title") or ""), 21, HEAD, True, 0)])
        _rect(s, 0.84, 1.24, 1.5, 0.035, ACCENT)
        bullets = item.get("bullets") or []
        lines = [("▪  " + _clean(b), 15, TEXT, False, 16) for b in bullets[:8]]
        if lines:
            _text(s, 0.95, 1.62, 11.5, 5.2, lines)
        _text(s, 0.84, 7.04, 9.0, 0.38, [(footer, 9, GRAY, False, 0)])
        _text(s, 12.0, 7.04, 0.9, 0.38, [(str(idx), 11, GRAY, False, 0)],
              align=PP_ALIGN.RIGHT)

    # ── 结尾页
    s = prs.slides.add_slide(blank)
    _rect(s, 0, 0, 13.333, 7.5, DEEP)
    _rect(s, 10.3, -1.9, 5.0, 5.0, DEEP2, MSO_SHAPE.OVAL)
    _rect(s, -1.3, 5.3, 3.2, 3.2, DEEP2, MSO_SHAPE.OVAL)
    _text(s, 0, 2.95, 13.333, 1.1, [("谢谢观看", 40, WHITE, True, 0)],
          align=PP_ALIGN.CENTER)
    _text(s, 0, 4.35, 13.333, 0.55, [(footer, 12, PALE, False, 0)],
          align=PP_ALIGN.CENTER)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PPT 模板填充（R8：上传模板自动填写）
# 模板占位符协议：封面页 {{title}}/{{subtitle}}/{{date}}/{{project}}；
# 内容页 {{slide_title}} + {{bullets}}（bullets 文本框按其首段格式逐条填充）；
# 结尾页（可选）{{closing}}/{{date}}/{{project}}。填充后删除模板原页，仅保留成品页。

def _ppt_clean(v) -> str:
    return str(v).replace("**", "").replace("`", "").strip()


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
    return "\n".join(parts)


_COPYABLE_REL_TYPES = (
    "image", "video", "media", "audio", "chart", "oleObject", "package", "hyperlink",
)


def _copy_slide(prs, source):
    """复制模板页（形状深拷贝 + 资源类关系重映射 rId）。

    只复制资源关系（图片/媒体/图表/外部链接等）；notesSlide/slide 等结构关系
    不复制——notesSlide 只能属于一页，多页共享会被 PowerPoint 判定为损坏文件。"""
    import copy as _copy

    dest = prs.slides.add_slide(source.slide_layout)
    for shp in list(dest.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map: dict[str, str] = {}
    for rId, rel in source.part.rels.items():
        base = rel.reltype.rsplit("/", 1)[-1]
        if base not in _COPYABLE_REL_TYPES:
            continue
        if rel.is_external:
            new_rid = dest.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            new_rid = dest.part.relate_to(rel.target_part, rel.reltype)
        if new_rid != rId:
            rid_map[rId] = new_rid
    for shp in source.shapes:
        el = _copy.deepcopy(shp._element)
        if rid_map:
            for node in el.iter():
                for attr, val in list(node.attrib.items()):
                    if val in rid_map and attr.split("}")[-1] in ("embed", "link", "id"):
                        node.attrib[attr] = rid_map[val]
        dest.shapes._spTree.append(el)
    return dest


def _remove_slide(prs, index: int):
    id_lst = prs.slides._sldIdLst
    sld = list(id_lst)[index]
    prs.part.drop_rel(sld.get(qn("r:id")))
    id_lst.remove(sld)


def _copy_font(src_run, dst_run):
    f, g = src_run.font, dst_run.font
    if f.size is not None:
        g.size = f.size
    if f.bold is not None:
        g.bold = f.bold
    try:
        if f.color and f.color.type is not None:
            g.color.rgb = f.color.rgb
    except Exception:
        pass
    if f.name:
        g.name = f.name
        _apply_ea_font(dst_run._r.get_or_add_rPr(), f.name)


def _apply_ea_font(rPr, name: str):
    """a:ea 必须紧跟 a:latin（Word/PowerPoint 对 rPr 子元素顺序严格校验）。"""
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
        latin = rPr.find(qn("a:latin"))
        if latin is not None:
            latin.addnext(ea)
    ea.set("typeface", name)


def _apply_map(slide, mapping: dict[str, str]) -> set[str]:
    """替换页内 {{key}} 占位符（保留首 run 格式），返回命中的 key。"""
    hit: set[str] = set()
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            full = "".join(r.text for r in para.runs)
            for key in list(mapping.keys()):
                if key in full:
                    new_text = full.replace(key, mapping[key])
                    if para.runs:
                        first = para.runs[0]
                        for r in para.runs[1:]:
                            r._r.getparent().remove(r._r)
                        first.text = new_text
                    else:
                        r = para.add_run()
                        r.text = new_text
                    hit.add(key)
                    full = new_text
    return hit


def _fill_bullets(shape, bullets: list[str]):
    """{{bullets}} 文本框：清空后按首段格式逐条填充。"""
    tf = shape.text_frame
    template_run = tf.paragraphs[0].runs[0] if tf.paragraphs[0].runs else None
    tx_body = tf._txBody
    for p in list(tx_body.findall(qn("a:p")))[1:]:
        tx_body.remove(p)
    first_para = tf.paragraphs[0]
    for r in list(first_para.runs):
        r._r.getparent().remove(r._r)
    for i, b in enumerate(bullets[:8]):
        para = first_para if i == 0 else tf.add_paragraph()
        run = para.add_run()
        run.text = _ppt_clean(b)
        if template_run is not None:
            _copy_font(template_run, run)


def _build_pptx_from_template(tmpl: bytes, outline: dict, project_name: str = "") -> bytes:
    """按占位符协议填充用户上传的 pptx 模板。"""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(tmpl))
    slides = list(prs.slides)
    now = datetime.now().strftime("%Y-%m-%d")
    cover_map = {
        "{{title}}": _ppt_clean(outline.get("title") or "主题分析报告"),
        "{{subtitle}}": _ppt_clean(outline.get("subtitle") or ""),
        "{{date}}": now,
        "{{project}}": project_name or "—",
    }
    closing_map = {"{{closing}}": "谢谢观看", "{{date}}": now,
                   "{{project}}": project_name or "—"}

    cover_idx = content_idx = closing_idx = None
    for i, s in enumerate(slides):
        text = _slide_text(s)
        if cover_idx is None and "{{title}}" in text:
            cover_idx = i
        if content_idx is None and "{{slide_title}}" in text:
            content_idx = i
        if closing_idx is None and "{{closing}}" in text:
            closing_idx = i
    if cover_idx is None or content_idx is None:
        raise APIError("模板缺少必需占位符：封面页须含 {{title}}，内容页须含 {{slide_title}}",
                       code="PPT_BAD_TEMPLATE", http_status=400)

    # 定位内容页的 {{bullets}} 文本框（缺失时回退到正文区最大文本框）
    content_src = slides[content_idx]
    bullets_shape = None
    fallback_shape = None
    for shape in content_src.shapes:
        if not shape.has_text_frame:
            continue
        if "{{bullets}}" in shape.text_frame.text:
            bullets_shape = shape
            break
        if "{{slide_title}}" not in shape.text_frame.text and shape.text_frame.text.strip():
            if fallback_shape is None or len(shape.text_frame.text) > len(fallback_shape.text_frame.text):
                fallback_shape = shape
    if bullets_shape is None:
        bullets_shape = fallback_shape
    if bullets_shape is None:
        raise APIError("内容页缺少 {{bullets}} 占位文本框",
                       code="PPT_BAD_TEMPLATE", http_status=400)

    # 1) 复制封面并填充（成品页一律为复制页，原模板页最后统一删除）
    cover_dest = _copy_slide(prs, slides[cover_idx])
    _apply_map(cover_dest, cover_map)

    # 2) 为每条大纲复制一张内容页并填充
    for item in outline["slides"]:
        page = _copy_slide(prs, content_src)
        _apply_map(page, {"{{slide_title}}": _ppt_clean(item.get("title") or "")})
        target = _find_shape_by_id(page, bullets_shape.shape_id)
        if target is not None:
            _fill_bullets(target, item.get("bullets") or [])

    # 3) 复制结尾页（追加在最后）
    if closing_idx is not None:
        closing_dest = _copy_slide(prs, slides[closing_idx])
        _apply_map(closing_dest, closing_map)

    # 4) 删除原模板页（成品 = 封面 + N 张内容页 + 结尾）
    for i in range(len(slides) - 1, -1, -1):
        _remove_slide(prs, i)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _find_shape_by_id(slide, shape_id: int):
    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape
    return None


SAMPLE_TEMPLATE_KEY = "ppt-templates/_sample.pptx"


def _build_sample_template() -> bytes:
    """内置示例模板（含占位符协议说明页，说明页放最后、生成时自动删除）。"""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Inches, Pt

    DEEP = RGBColor(0x1B, 0x3C, 0x6E)
    DEEP2 = RGBColor(0x25, 0x50, 0x8F)
    ACCENT = RGBColor(0x5B, 0x8D, 0xEF)
    HEAD = RGBColor(0x1F, 0x4E, 0x79)
    TEXT = RGBColor(0x33, 0x3A, 0x45)
    GRAY = RGBColor(0x8A, 0x94, 0xA6)
    PALE = RGBColor(0xBF, 0xD4, 0xF5)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    FONT = "微软雅黑"

    def run_style(run, size, color, bold=False):
        f = run.font
        f.size = Pt(size); f.color.rgb = color; f.bold = bold; f.name = FONT
        _apply_ea_font(run._r.get_or_add_rPr(), FONT)

    def rect(slide, x, y, w, h, color, oval=False):
        sp = slide.shapes.add_shape(MSO_SHAPE.OVAL if oval else MSO_SHAPE.RECTANGLE,
                                    Inches(x), Inches(y), Inches(w), Inches(h))
        sp.fill.solid(); sp.fill.fore_color.rgb = color
        sp.line.fill.background(); sp.shadow.inherit = False
        return sp

    def text(slide, x, y, w, h, content, size, color, bold=False,
             align=PP_ALIGN.LEFT, line_spacing=1.2):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for i, line in enumerate(content.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line_spacing
            p.space_after = Pt(4)
            r = p.add_run(); r.text = line
            run_style(r, size, color, bold)
        return tb

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1) 封面
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, DEEP)
    rect(s, 9.9, -2.3, 5.8, 5.8, DEEP2, oval=True)
    rect(s, -1.5, 5.5, 3.6, 3.6, DEEP2, oval=True)
    rect(s, 0.95, 2.62, 0.10, 1.62, ACCENT)
    text(s, 1.35, 2.42, 10.7, 2.0, "{{title}}", 33, WHITE, True)
    text(s, 1.35, 4.62, 10.6, 0.8, "{{subtitle}}", 16, PALE)
    text(s, 1.35, 6.74, 10.5, 0.5, "{{project}} · {{date}}", 11, PALE)

    # 2) 内容页（{{bullets}} 文本框预留足够高度）
    s = prs.slides.add_slide(blank)
    rect(s, 0.55, 0.50, 0.10, 0.48, ACCENT)
    text(s, 0.82, 0.38, 11.8, 0.75, "{{slide_title}}", 21, HEAD, True)
    rect(s, 0.84, 1.24, 1.5, 0.035, ACCENT)
    text(s, 0.95, 1.62, 11.5, 5.2, "{{bullets}}", 15, TEXT)

    # 3) 结尾页
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, DEEP)
    rect(s, 10.3, -1.9, 5.0, 5.0, DEEP2, oval=True)
    rect(s, -1.3, 5.3, 3.2, 3.2, DEEP2, oval=True)
    text(s, 0, 2.95, 13.333, 1.1, "{{closing}}", 40, WHITE, True, align=PP_ALIGN.CENTER)
    text(s, 0, 4.35, 13.333, 0.55, "{{project}} · {{date}}", 12, PALE, align=PP_ALIGN.CENTER)

    # 4) 说明页（放最后；生成时作为模板原页自动删除，不会出现在成品中）
    s = prs.slides.add_slide(blank)
    text(s, 0.8, 0.5, 11.7, 0.8, "PPT 模板使用说明（本页为说明页，生成时自动删除）", 20, HEAD, True)
    guide = (
        "把本文件第 1~3 页的样式换成你的企业模板即可：保留占位符文本，其余设计（配色/字体/图片/Logo/版式）随意发挥，生成时全部保留。\n"
        "占位符协议（占位符必须完整写在同一段落内）：\n"
        "  · 第 1 页 封面：{{title}} 主标题（必需）｜{{subtitle}} 副标题｜{{project}} 项目名｜{{date}} 生成日期\n"
        "  · 第 2 页 内容页：{{slide_title}} 页标题（必需）｜{{bullets}} 要点文本框（按大纲逐条填入，每页最多 8 条，请预留足够高度；该文本框的字体字号即成品样式）\n"
        "  · 第 3 页 结尾页（可选）：{{closing}} 结束语｜{{project}}｜{{date}}\n"
        "注意事项：\n"
        "  · 内容页只需制作 1 张，系统按大纲页数自动复制；多余页面会被删除，成品 = 封面 + N 页内容 + 结尾\n"
        "  · 占位符可以与其他文字混排（如落款写 {{project}} · {{date}}）\n"
        "  · 缺少 {{title}} 或 {{slide_title}} 时生成会报错并提示缺哪一项\n"
        "  · 文件须为 .pptx 格式、不超过 20MB；改完直接在「自定义模板」处上传即可"
    )
    text(s, 0.8, 1.5, 11.7, 5.6, guide, 13, TEXT, line_spacing=1.25)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _ensure_sample_template() -> str:
    """示例模板已在 MinIO 则返回预签名 URL，否则生成并上传。"""
    mc = get_minio_client()
    if mc.stat_object(settings.MINIO_BUCKET_EXPORTS, SAMPLE_TEMPLATE_KEY) is None:
        mc.put_bytes(settings.MINIO_BUCKET_EXPORTS, SAMPLE_TEMPLATE_KEY,
                     _build_sample_template(),
                     content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation")
    return mc.get_presigned_url(settings.MINIO_BUCKET_EXPORTS, SAMPLE_TEMPLATE_KEY)


def _upload_export(project_id: int, filename: str, data: bytes, content_type: str) -> str:
    key = f"reports/{project_id}/{uuid.uuid4().hex[:8]}-{filename}"
    get_minio_client().put_bytes(settings.MINIO_BUCKET_EXPORTS, key, data,
                                 content_type=content_type)
    return get_minio_client().get_presigned_url(settings.MINIO_BUCKET_EXPORTS, key)


# ---------------------------------------------------------------- 端点

SUGGEST_SYSTEM = (
    "你是知识图谱分析顾问。基于给定的知识图谱材料（类分布、代表实体及其属性、关系谓词、文档清单），"
    "提出 6 个适合撰写业务分析报告的具体主题。要求："
    "1. 每个主题是一句话，20~35 字，面向业务读者（产品、风控、合规、投资分析等视角）；"
    "2. 主题必须能由材料中的事实支撑（须用到材料中出现的实体/关系/属性），不得超出材料范围编造；"
    "3. 六个主题之间分析角度差异化（如要素梳理、对比、风险、流程、合规、溯源等）；"
    '4. 只输出 JSON：{"topics": ["主题1", "主题2", ...]}，不要 markdown 代码块。'
)

_SUGGEST_CACHE: dict[int, tuple[float, list[str]]] = {}  # project_id -> (ts, topics)
_SUGGEST_TTL = 600  # 建议缓存 10 分钟（避免每次进页都等 LLM）


def _parse_topics(raw: str) -> list[str]:
    text = raw.strip()
    m = re.search(r"```(?:json)?\s*(.*?)```", text, re.S)
    if m:
        text = m.group(1).strip()
    start, end = text.find("{"), text.rfind("}")
    if start >= 0 and end > start:
        text = text[start:end + 1]
    data = json.loads(text)
    return [str(t).strip() for t in (data.get("topics") or []) if str(t).strip()][:8]


def _fallback_chips(ctx: dict) -> list[str]:
    """LLM 不可用时的兜底：图谱高频类/代表实体/业务谓词。"""
    chips: list[str] = []
    for c in ctx["classes"][:5]:
        if c["class"]:
            chips.append(c["class"])
    for c in ctx["classes"][:3]:
        for e in c["entities"][:2]:
            if e["label"] and e["label"] not in chips:
                chips.append(e["label"])
    for r in ctx["top_relations"][:3]:
        if r["predicate"] != "instance_of" and r["predicate"] not in chips:
            chips.append(r["predicate"])
    return chips[:12]


@router.get("/{project_id}/reports/suggest", dependencies=[Depends(require_any_module("report", "ppt"))])
def suggest_topics(project_id: int, db: Session = Depends(get_db),
                   user: User = Depends(get_current_user)):
    """主题建议：LLM 基于图谱材料生成具体分析主题（缓存 10 分钟；失败回退知识锚点）。"""
    _ensure_qa_access(db, project_id, user)
    ctx = _collect_context(db, project_id)
    if ctx["entity_total"] == 0:
        return {"chips": [], "source": "empty"}

    cached = _SUGGEST_CACHE.get(project_id)
    if cached and time.time() - cached[0] < _SUGGEST_TTL:
        return {"chips": cached[1], "source": "cache"}

    try:
        cfg = build_legacy_llm_config(db, project_id)
        llm = LLMClient(api_key=cfg.get("api_key"), base_url=cfg.get("base_url"),
                        model=cfg.get("model"))
        result = llm.call_llm_text(SUGGEST_SYSTEM,
                                   f"知识图谱材料：\n{_context_text(ctx)}", timeout=60)
        content = result.get("content") if isinstance(result, dict) else str(result)
        if not content or not content.strip():
            raise APIError("模型未返回主题建议", code="SUGGEST_EMPTY", http_status=502)
        topics = _parse_topics(content)
        if not topics:
            raise APIError("主题建议解析失败", code="SUGGEST_BAD_FORMAT", http_status=502)
    except Exception as e:  # LLM 不可用/解析失败 → 兜底锚点，页面不阻塞
        logger.warning(f"[reports] LLM 主题建议失败，回退知识锚点：{e}")
        return {"chips": _fallback_chips(ctx), "source": "fallback"}

    _SUGGEST_CACHE[project_id] = (time.time(), topics)
    return {"chips": topics, "source": "llm"}


@router.post("/{project_id}/reports/generate", dependencies=[Depends(require_any_module("report", "ppt"))])
def generate_report(
    project_id: int,
    body: ReportRequest,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """主题分析报告/PPT（R8）：围绕 topic 检索图谱证据 → LLM 生成面向业务读者的分析。
    模块按 kind 分流：report→[M:report]，ppt→[M:ppt]。"""
    ensure_module(user, db, "report" if body.kind == "report" else "ppt")
    proj = _ensure_qa_access(db, project_id, user)
    t0 = time.time()

    topic = body.topic.strip()
    ctx = _collect_context(db, project_id, topic=topic)
    if ctx["entity_total"] == 0:
        raise APIError("该项目尚无抽取完成的图谱数据，请先完成实例抽取",
                       code="REPORT_NO_DATA", http_status=400)

    cfg = build_legacy_llm_config(db, project_id)
    llm = LLMClient(api_key=cfg.get("api_key"), base_url=cfg.get("base_url"),
                    model=cfg.get("model"))

    material = (f"项目名称：{proj.name}\n项目描述：{proj.description or '无'}\n\n"
                f"证据材料：\n{_context_text(ctx)}")

    downloads: dict[str, str] = {}
    stamp = datetime.now().strftime("%Y%m%d_%H%M")

    if body.kind == "report":
        content = _llm_call(llm, REPORT_SYSTEM, material)
        title_match = re.search(r"^#\s+(.+)$", content, re.M)
        title = (title_match.group(1).strip() if title_match else "本体分析报告")[:80]
        # 附录在 LLM 生成后确定性拼接：推理产物不进正文材料，避免冒充事实
        if body.include_inferred:
            try:
                rows = _inference_rows(db, project_id)
                if rows:
                    content = content.rstrip() + "\n\n---\n\n" + _inference_appendix_md(rows)
            except Exception as e:
                logger.warning(f"[reports] 推理附录生成失败，忽略：{e}")
        md_bytes = content.encode("utf-8")
        docx_bytes = _markdown_to_docx(content, title=title, project_name=proj.name,
                                       entity_total=ctx["entity_total"],
                                       relation_total=ctx["relation_total"])
        downloads["md"] = _upload_export(project_id, f"{title}_{stamp}.md", md_bytes,
                                         "text/markdown")
        downloads["docx"] = _upload_export(project_id, f"{title}_{stamp}.docx", docx_bytes,
                                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        result = {"kind": "report", "title": title, "content": content, "outline": None,
                  "downloads": downloads}
    else:
        raw = _llm_call(llm, PPT_SYSTEM, material + f"\n页数上限：{body.max_slides}")
        outline = _parse_ppt_outline(raw)
        # 附录页确定性追加（置于结尾内容之前，不参与 LLM 大纲）
        if body.include_inferred:
            try:
                rows = _inference_rows(db, project_id)
                if rows:
                    outline["slides"].append(_inference_appendix_slide(rows))
            except Exception as e:
                logger.warning(f"[reports] 推理附录页生成失败，忽略：{e}")
        if body.use_template:
            if get_minio_client().stat_object(settings.MINIO_BUCKET_EXPORTS,
                                              _template_key(project_id)) is None:
                raise APIError("尚未上传 PPT 模板，请先上传或关闭「使用模板」",
                               code="PPT_TEMPLATE_MISSING", http_status=400)
            tmpl = get_minio_client().get_bytes(settings.MINIO_BUCKET_EXPORTS,
                                                _template_key(project_id))
            pptx_bytes = _build_pptx_from_template(tmpl, outline, project_name=proj.name)
            used_template = True
        else:
            pptx_bytes = _build_pptx(outline, project_name=proj.name)
            used_template = False
        safe_title = re.sub(r'[\\/:*?"<>|]', '_', str(outline.get("title") or "本体演示"))[:50]
        downloads["pptx"] = _upload_export(project_id, f"{safe_title}_{stamp}.pptx", pptx_bytes,
                                           "application/vnd.openxmlformats-officedocument.presentationml.presentation")
        result = {"kind": "ppt", "title": outline.get("title") or "本体演示",
                  "content": None, "outline": outline, "downloads": downloads,
                  "used_template": used_template}

    result["latency_ms"] = int((time.time() - t0) * 1000)
    log_action(db, user.id, "report.generate", resource_type="project",
               resource_id=str(project_id),
               detail={"kind": body.kind, "latency_ms": result["latency_ms"],
                       "include_inferred": body.include_inferred})
    db.commit()
    return result

